"""Offline document readers. No document text, filenames or OCR errors are logged."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from html import escape
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import threading
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from markus_dispatcher.core.limited_process import OCR_MEMORY_BYTES, ResourceLimitError, run_limited


MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_XML_BYTES = 20 * 1024 * 1024
MAX_CELLS = 100_000
OCR_LANGUAGES = ('pol', 'eng', 'deu', 'fra', 'ita')
_OCR_SLOT = threading.BoundedSemaphore(1)
WORD_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
FILE_FILTER = ('Dokumenty (*.pdf *.docx *.doc *.xlsx *.xls *.xlsm *.csv *.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp);;'
               'PDF (*.pdf);;Word (*.docx *.doc);;Excel (*.xlsx *.xls *.xlsm *.csv);;'
               'Obrazy (*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp)')


@dataclass(frozen=True)
class PreviewSection:
    title: str
    html: str


def check_file(path: Path) -> None:
    if not path.is_file():
        raise ValueError('Nie znaleziono wybranego pliku.')
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError('Plik jest większy niż 50 MB. Podziel go lokalnie na mniejsze części.')


def _xml(archive: ZipFile, name: str):
    member = archive.getinfo(name)
    if member.file_size > MAX_XML_BYTES:
        raise ValueError('Dokument zawiera zbyt duży element do podglądu.')
    data = archive.read(name)
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        raise ValueError('Dokument zawiera nieobsługiwaną deklarację XML.')
    return ET.fromstring(data)


def _paragraph(node) -> str:
    parts = []
    for item in node.iter():
        if item.tag == WORD_NS + 't':
            parts.append(escape(item.text or ''))
        elif item.tag in (WORD_NS + 'br', WORD_NS + 'cr'):
            parts.append('<br>')
        elif item.tag == WORD_NS + 'tab':
            parts.append(' &nbsp; ')
    return ''.join(parts)


def _word_blocks(parent) -> str:
    parts = []
    for item in parent:
        if item.tag == WORD_NS + 'p':
            parts.append('<p>' + _paragraph(item) + '</p>')
        elif item.tag == WORD_NS + 'tbl':
            rows = []
            for row in item.findall(WORD_NS + 'tr'):
                cells = []
                for cell in row.findall(WORD_NS + 'tc'):
                    span = cell.find(WORD_NS + 'tcPr/' + WORD_NS + 'gridSpan')
                    colspan = min(100, int(span.get(WORD_NS + 'val', '1'))) if span is not None else 1
                    cells.append(f'<td colspan="{colspan}">' + _word_blocks(cell) + '</td>')
                rows.append('<tr>' + ''.join(cells) + '</tr>')
            parts.append('<table border="1" cellspacing="0" cellpadding="8" width="100%">' + ''.join(rows) + '</table>')
    return ''.join(parts)


def read_docx(path: Path) -> list[PreviewSection]:
    check_file(path)
    with ZipFile(path) as archive:
        root = _xml(archive, 'word/document.xml')
        body = root.find(WORD_NS + 'body')
        sections = [PreviewSection('Dokument', _word_blocks(body))] if body is not None else []
        # Headers/footers may contain the POD recipient; do not silently omit them.
        for name in sorted(archive.namelist()):
            if re.fullmatch(r'word/(header|footer)\d+\.xml', name):
                sections.append(PreviewSection('Nagłówek / stopka ' + str(len(sections)), _word_blocks(_xml(archive, name))))
        return sections


def _table(rows) -> str:
    parts = []
    for row in rows:
        cells = [escape(str(value) if value is not None else '').replace('\n', '<br>') for value in row]
        parts.append('<tr>' + ''.join('<td>' + cell + '</td>' for cell in cells) + '</tr>')
    return '<table border="1" cellspacing="0" cellpadding="7">' + ''.join(parts) + '</table>'


def read_spreadsheet(path: Path) -> list[PreviewSection]:
    check_file(path)
    if path.suffix.lower() == '.csv':
        data = path.read_bytes()
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('cp1250')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=';,\t|')
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(text.splitlines(), dialect))
        if sum(len(row) for row in rows) > MAX_CELLS:
            raise ValueError('Arkusz przekracza limit 100 000 komórek podglądu.')
        return [PreviewSection('CSV', _table(rows))]
    if path.suffix.lower() == '.xls':
        try:
            import xlrd
        except ImportError:
            raise ValueError('Obsługa XLS wymaga lokalnego pakietu xlrd. Możesz też zapisać plik jako XLSX lub PDF.') from None
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            result = []
            for sheet in workbook.sheets():
                if sheet.nrows * sheet.ncols > MAX_CELLS:
                    raise ValueError('Arkusz przekracza limit 100 000 komórek podglądu.')
                rows = []
                for r in range(sheet.nrows):
                    row = []
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode).strftime('%d.%m.%Y %H:%M')
                        row.append(value)
                    rows.append(row)
                result.append(PreviewSection(sheet.name, _table(rows)))
            return result
        finally:
            workbook.release_resources()
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError('Obsługa XLSX wymaga lokalnego pakietu openpyxl.') from None
    with ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Arkusz po rozpakowaniu przekracza limit podglądu.')
    workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    try:
        result = []
        for sheet in workbook.worksheets:
            if sheet.max_row * sheet.max_column > MAX_CELLS:
                raise ValueError('Arkusz przekracza limit 100 000 komórek podglądu.')
            spans, covered = {}, set()
            for region in sheet.merged_cells.ranges:
                spans[(region.min_row, region.min_col)] = (region.max_row-region.min_row+1, region.max_col-region.min_col+1)
                covered.update((r, c) for r in range(region.min_row, region.max_row+1)
                               for c in range(region.min_col, region.max_col+1) if (r, c) != (region.min_row, region.min_col))
            rows = []
            for row in sheet.iter_rows():
                cells = []
                for cell in row:
                    if (cell.row, cell.column) in covered:
                        continue
                    rs, cs = spans.get((cell.row, cell.column), (1, 1))
                    value = cell.value
                    if isinstance(value, datetime):
                        value = value.strftime('%d.%m.%Y %H:%M')
                    text = escape(str(value) if value is not None else '').replace('\n', '<br>')
                    cells.append(f'<td rowspan="{rs}" colspan="{cs}">{text}</td>')
                rows.append('<tr>' + ''.join(cells) + '</tr>')
            result.append(PreviewSection(sheet.title, '<table border="1" cellspacing="0" cellpadding="7">' + ''.join(rows) + '</table>'))
        return result
    finally:
        workbook.close()


def tesseract_path() -> str | None:
    private_roots = [Path(getattr(sys, '_MEIPASS', Path(sys.executable).parent))/'tesseract',
                     Path(__file__).resolve().parents[2]/'work'/'local-tools'/'tesseract']
    for root in private_roots:
        candidate = root/'tesseract.exe'
        if candidate.is_file():
            return str(candidate)
    found = shutil.which('tesseract')
    if found:
        return found
    for root in (os.environ.get('ProgramFiles'), os.environ.get('ProgramFiles(x86)')):
        if root:
            candidate = Path(root) / 'Tesseract-OCR' / 'tesseract.exe'
            if candidate.is_file():
                return str(candidate)
    return None


def run_local_ocr(png: bytes) -> str:
    engine = tesseract_path()
    if not engine:
        raise ValueError('Nie znaleziono lokalnego Tesseract OCR. Zainstaluj go lub wpisz treść zaznaczenia ręcznie.')
    if not _OCR_SLOT.acquire(blocking=False):
        raise ValueError('Inny odczyt OCR jest w toku. Poczekaj na jego zakończenie, aby zachować wspólny limit 2 GB.')
    env = {**os.environ, 'OMP_THREAD_LIMIT': '2', 'OMP_NUM_THREADS': '2', 'OMP_DYNAMIC': 'FALSE'}
    data_dir = Path(engine).parent/'tessdata'
    data_args = ['--tessdata-dir', str(data_dir)] if data_dir.is_dir() else []
    try:
        available = run_limited([engine, *data_args, '--list-langs'], timeout=10, env=env, memory_bytes=OCR_MEMORY_BYTES)
        if available.returncode:
            raise ValueError('Nie udało się odczytać lokalnych danych językowych OCR.')
        langs = set(available.stdout.decode('utf-8', errors='replace').splitlines())
        chosen = [code for code in OCR_LANGUAGES if code in langs]
        if not chosen:
            raise ValueError('Tesseract nie ma danych językowych pol, eng, deu, fra ani ita.')
        # No files with document contents are needed; stdin/stdout stay local.
        result = run_limited([engine, 'stdin', 'stdout', *data_args, '-l', '+'.join(chosen), '--psm', '6'],
                             input=png, timeout=45, env=env, memory_bytes=OCR_MEMORY_BYTES)
        if result.returncode:
            raise ValueError('Odczyt OCR przerwano z powodu błędu lub limitu pamięci 2 GB. Zaznacz mniejszy obszar albo wpisz treść ręcznie.')
        return result.stdout.decode('utf-8', errors='replace').strip()
    except subprocess.TimeoutExpired:
        raise ValueError('Przerwano OCR po przekroczeniu limitu czasu. Zaznacz mniejszy obszar.') from None
    except ResourceLimitError as error:
        raise ValueError(str(error)) from None
    finally:
        _OCR_SLOT.release()


def convert_doc(path: Path, target: Path) -> Path:
    """Legacy DOC through a dedicated, local LibreOffice process; no shell interpolation."""
    check_file(path)
    engine = shutil.which('soffice')
    if not engine:
        for base in (os.environ.get('ProgramFiles'), os.environ.get('ProgramFiles(x86)')):
            candidate = Path(base or '') / 'LibreOffice' / 'program' / 'soffice.exe'
            if candidate.is_file():
                engine = str(candidate)
                break
    if not engine:
        raise ValueError('Stary format DOC wymaga lokalnego LibreOffice. Zapisz dokument jako DOCX lub PDF albo zainstaluj LibreOffice.')
    # Isolated profile disables macros, link updates and Java. Source copied under
    # a fixed basename, so control characters/options in filenames cannot run code.
    import shutil as files
    with tempfile.TemporaryDirectory(prefix='spedi-doc-', ignore_cleanup_errors=True) as temporary:
        work = Path(temporary)
        source = work / 'document.doc'
        files.copyfile(path, source)
        profile = work / 'profile'
        user = profile / 'user'
        user.mkdir(parents=True)
        (user / 'registrymodifications.xcu').write_text(
            '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop></item>'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="SecureURL" oor:op="fuse"><value/></prop></item>'
            '<item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item>'
            '<item oor:path="/org.openoffice.Office.Java/VirtualMachine"><prop oor:name="Enable" oor:op="fuse"><value>false</value></prop></item></oor:items>', encoding='utf-8')
        try:
            result = subprocess.run([engine, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--norestore',
                                     '--convert-to', 'pdf', '--outdir', str(work), str(source)],
                                    capture_output=True, timeout=60,
                                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        except subprocess.TimeoutExpired:
            raise ValueError('Przerwano konwersję DOC po 60 sekundach. Zapisz dokument lokalnie jako PDF.') from None
        output = work / 'document.pdf'
        if result.returncode or not output.is_file():
            raise ValueError('Nie udało się przekonwertować DOC lokalnie. Zapisz go jako PDF lub DOCX.')
        target.mkdir(parents=True, exist_ok=True)
        return Path(files.copyfile(output, target / 'document.pdf'))


def parse_import_datetime(value: str) -> str:
    match = re.search(r'(?<!\d)(\d{1,2}[./-]\d{1,2}[./-]\d{4}|\d{4}-\d{2}-\d{2})\s*(?:r\.?\s*)?(?:[T,]|godz\.?)?\s*(\d{1,2}[:.]\d{2})(?!\d)', value)
    if not match:
        raise ValueError('Wskaż jedną pełną datę z godziną, np. 26.10.2026 08:00. Zakres godzin wpisz w uwagach.')
    if len(re.findall(r'\d{1,2}:\d{2}', value)) > 1:
        raise ValueError('Zaznaczono kilka godzin. Wybierz jedną i zachowaj pełne okno czasowe w uwagach.')
    date, time = match.groups()
    for fmt in ('%d.%m.%Y', '%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(date + ' ' + time.replace('.', ':'), fmt + ' %H:%M').strftime('%d.%m.%Y %H:%M')
        except ValueError:
            continue
    raise ValueError('Nieprawidłowa data lub godzina.')
