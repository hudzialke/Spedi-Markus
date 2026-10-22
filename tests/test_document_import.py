"""Synthetic fixtures only: never open dispatcher user data or real documents."""
from datetime import datetime
import json
import os
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest

from markus_dispatcher.core import document_import as readers
from markus_dispatcher.core.models import JobData
from markus_dispatcher.data import db


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    monkeypatch.setenv('TEMP', str(tmp_path))
    monkeypatch.setenv('TMP', str(tmp_path))
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.fixture
def local_db(tmp_path, monkeypatch):
    for key, value in (('DATA_DIR', tmp_path), ('DB_PATH', tmp_path/'test.db'),
                       ('ARCHIVE_ROOT', tmp_path/'archive'), ('ROZPISKA_DB_CANDIDATE', tmp_path/'none.db'),
                       ('FOOTER_PATH', tmp_path/'footer.txt')):
        monkeypatch.setattr(db.config, key, value)
    db.init_db()
    return db.add_driver('Synthetic driver')


@pytest.fixture
def columns_pdf(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    path = tmp_path/'synthetic-columns.pdf'
    canvas = Canvas(str(path), pagesize=(600, 800))
    # Deliberately interleave left/right text in the PDF content stream.
    for x, y, text in ((50,700,'LOAD Demo A'), (340,700,'UNLOAD Demo B'),
                       (50,680,'Test Street 12'), (340,680,'Other Street 45'),
                       (50,400,'POD: PDF within 24 hours')):
        canvas.drawString(x, y, text)
    canvas.showPage()
    canvas.drawString(50, 700, 'Page two conditions')
    canvas.drawString(50, 680, '00-001 Warsaw, PL; 12/3 (B)')
    canvas.save()
    return path


def test_pdf_rectangles_keep_columns_separate_even_with_interleaved_text(app, columns_pdf):
    from PySide6.QtCore import QRectF
    from PySide6.QtPdf import QPdfDocument
    from markus_dispatcher.ui.document_import_dialog import pdf_text_in_rect
    document = QPdfDocument()
    assert document.load(str(columns_pdf)) == QPdfDocument.Error.None_
    assert pdf_text_in_rect(document, 0, QRectF(40, 80, 220, 80)) == 'LOAD Demo A\nTest Street 12'
    assert pdf_text_in_rect(document, 0, QRectF(330, 80, 220, 80)) == 'UNLOAD Demo B\nOther Street 45'
    assert pdf_text_in_rect(document, 0, QRectF(40, 350, 230, 80)) == 'POD: PDF within 24 hours'
    assert pdf_text_in_rect(document, 1, QRectF(40, 105, 300, 30)) == '00-001 Warsaw, PL; 12/3 (B)'
    document.close()


def test_pdf_mapping_clears_selection_on_page_change_and_retains_source(app, columns_pdf):
    from PySide6.QtCore import QRectF
    from markus_dispatcher.ui.document_import_dialog import DocumentImportDialog, FIELDS
    dialog = DocumentImportDialog()
    dialog.open_path(columns_pdf)
    assert dialog.pdf.pageCount() == 2
    dialog._activate(next(i for i,(key,_) in enumerate(FIELDS) if key == 'load_address'))
    dialog._region(QRectF(40/600, 80/800, 220/600, 80/800))
    dialog.assign_selection()
    assert dialog.values['load_address'] == 'LOAD Demo A\nTest Street 12'
    assert dialog.sources['load_address']['page'] == 1
    assert 'region' in dialog.sources['load_address']
    dialog.page_combo.setCurrentIndex(1)
    assert not dialog.selection.toPlainText()
    assert dialog._crop.isNull()
    assert not dialog.ocr_btn.isEnabled()
    dialog.reject()


def test_docx_keeps_table_cells_and_escapes_content(tmp_path):
    path = tmp_path/'synthetic.docx'
    xml = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:tbl><w:tr><w:tc><w:p><w:r><w:t>LOAD &lt;demo&gt;</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>UNLOAD</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>'
    with ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', xml)
        archive.writestr('word/footer1.xml', '<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>POD requirement</w:t></w:r></w:p></w:ftr>')
    sections = readers.read_docx(path)
    assert sections[0].html.count('<td ') == 2
    assert 'LOAD &lt;demo&gt;' in sections[0].html
    assert 'POD requirement' in sections[1].html


def test_docx_external_relationships_are_not_loaded(tmp_path, monkeypatch):
    path = tmp_path/'synthetic.docx'
    with ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Local text</w:t></w:r></w:p></w:body></w:document>')
        archive.writestr('word/_rels/document.xml.rels', '<Relationships><Relationship TargetMode="External" Target="https://example.invalid/image"/></Relationships>')
    import socket
    monkeypatch.setattr(socket, 'create_connection', lambda *a, **k: pytest.fail('No network'))
    assert 'Local text' in readers.read_docx(path)[0].html


def test_xlsx_preserves_merged_cells_dates_and_multiple_sheets(tmp_path):
    from openpyxl import Workbook
    path = tmp_path/'synthetic.xlsx'
    book = Workbook()
    sheet = book.active
    sheet['A1'] = 'LOAD'
    sheet.merge_cells('A1:B1')
    sheet['C1'] = 'UNLOAD'
    sheet['A2'] = datetime(2026,10,26,8)
    sheet['C2'] = 'Test <street>'
    book.create_sheet('POD')['A1'] = 'PDF in 24h'
    book.save(path)
    sections = readers.read_spreadsheet(path)
    assert len(sections) == 2
    assert 'colspan="2"' in sections[0].html
    assert '26.10.2026 08:00' in sections[0].html
    assert 'Test &lt;street&gt;' in sections[0].html
    assert 'PDF in 24h' in sections[1].html


def test_csv_semicolon_preserves_columns(tmp_path):
    path = tmp_path/'synthetic.csv'
    path.write_text('Load;Unload\nDemo A;Demo B', encoding='utf-8-sig')
    html = readers.read_spreadsheet(path)[0].html
    assert '<td>Demo A</td><td>Demo B</td>' in html


def test_spreadsheet_dialog_worker_and_text_assignment(app, tmp_path):
    from markus_dispatcher.ui.document_import_dialog import DocumentImportDialog
    path = tmp_path/'synthetic.csv'
    path.write_text('Load;Unload\nDemo A;Demo B', encoding='utf-8')
    dialog = DocumentImportDialog()
    dialog.open_path(path)
    assert dialog._reader.wait(5000)
    app.processEvents()
    assert dialog._has_document
    assert dialog.preview_stack.currentWidget() == dialog.text_preview
    cursor = dialog.text_preview.document().find('Demo A')
    dialog.text_preview.setTextCursor(cursor)
    dialog.assign_selection()
    assert dialog.values['ref_numbers'] == 'Demo A'
    assert dialog.sources['ref_numbers']['section'] == 'CSV'
    dialog.reject()


@pytest.mark.parametrize('value,expected', [('Termin: 26.10.2026 8:00','26.10.2026 08:00'),
                                          ('2026-10-26T08:00','26.10.2026 08:00'),
                                          ('26/10/2026 godz. 08:00','26.10.2026 08:00')])
def test_datetime_normalization(value, expected):
    assert readers.parse_import_datetime(value) == expected


@pytest.mark.parametrize('value', ['26.10.2026', '31.02.2026 08:00', '26.10.2026 08:00-12:00', '26.10.2026 25:00'])
def test_dates_never_invent_time_or_accept_ranges(value):
    with pytest.raises(ValueError):
        readers.parse_import_datetime(value)


def test_local_ocr_uses_stdin_bounded_threads_timeout_and_no_shell(monkeypatch):
    calls = []
    monkeypatch.setattr(readers, 'tesseract_path', lambda: 'tesseract.exe')
    def run(args, **kwargs):
        calls.append((args, kwargs))
        return SimpleNamespace(returncode=0, stdout=b'List\neng\npol\ndeu\nfra\nita\n' if '--list-langs' in args else b'Demo OCR')
    monkeypatch.setattr(readers, 'run_limited', run)
    assert readers.run_local_ocr(b'synthetic png bytes') == 'Demo OCR'
    args, kwargs = calls[1]
    assert args[1:3] == ['stdin', 'stdout']
    assert kwargs['input'] == b'synthetic png bytes'
    assert kwargs['env']['OMP_THREAD_LIMIT'] == '2'
    assert kwargs['timeout'] == 45
    assert kwargs['memory_bytes'] == 2 * 1024**3
    assert args[args.index('-l')+1] == 'pol+eng+deu+fra+ita'
    assert not kwargs.get('shell')


def test_local_ocr_timeout_does_not_expose_raw_error(monkeypatch):
    monkeypatch.setattr(readers, 'tesseract_path', lambda: 'tesseract.exe')
    def fail(*a, **k):
        raise readers.subprocess.TimeoutExpired('private filename', 45, output=b'private OCR')
    monkeypatch.setattr(readers, 'run_limited', fail)
    with pytest.raises(ValueError) as error:
        readers.run_local_ocr(b'fake')
    assert 'private' not in str(error.value)


def test_ocr_does_not_start_second_concurrent_worker(monkeypatch):
    monkeypatch.setattr(readers, 'tesseract_path', lambda: 'tesseract.exe')
    monkeypatch.setattr(readers, 'run_limited', lambda *a, **k: pytest.fail('Second worker must not run'))
    assert readers._OCR_SLOT.acquire(blocking=False)
    try:
        with pytest.raises(ValueError, match='Inny odczyt'):
            readers.run_local_ocr(b'fake')
    finally:
        readers._OCR_SLOT.release()


@pytest.mark.skipif(os.name != 'nt', reason='Windows OCR memory limit')
def test_installed_ocr_reads_synthetic_french_and_italian_image():
    if not readers.tesseract_path():
        pytest.skip('Local OCR engine not installed')
    from io import BytesIO
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new('RGB', (1600, 340), 'white')
    font = ImageFont.truetype('C:/Windows/Fonts/segoeui.ttf', 40)
    drawing = ImageDraw.Draw(image)
    drawing.text((35, 35), 'POD: livraison sous 24 heures.', fill='black', font=font)
    drawing.text((35, 110), 'Consegna: inviare il documento in PDF.', fill='black', font=font)
    drawing.text((35, 185), 'Zaladunek: Warszawa. Delivery: Berlin.', fill='black', font=font)
    buffer = BytesIO()
    image.save(buffer, format='PNG')
    text = readers.run_local_ocr(buffer.getvalue()).casefold()
    assert 'livraison' in text
    assert 'consegna' in text
    assert 'pdf' in text


def test_import_saves_pod_sources_and_original_without_shared_client(app, local_db, tmp_path, monkeypatch, columns_pdf):
    from markus_dispatcher.ui.job_pane import JobPane
    import markus_dispatcher.ui.job_pane as module
    monkeypatch.setattr(module, 'ARCHIVE_ROOT', db.config.ARCHIVE_ROOT)
    monkeypatch.setattr(module, 'start_geocode', lambda *a: pytest.fail('No geocoding'))
    pane = JobPane()
    pane.new_job(driver_id=local_db.id)
    values = {'ref_numbers':'SYNTH-02', 'client_name':'Synthetic client', 'load_address':'Demo A 12',
              'unload_address':'Demo B 45', 'load_datetime':'26.10.2026 08:00',
              'unload_datetime':'27.10.2026 10:00', 'pod_requirements':'CMR as PDF in 24h', 'price':'650,00'}
    source = {'load_address': {'document':'synthetic.pdf', 'page':1, 'region':[0,0,.4,.4]}}
    pane._apply_document_import(values, source, columns_pdf)
    pane._save(show_message=False)
    job = db.get_job(pane._job_id)
    assert job.local_only and job.pod_requirements == values['pod_requirements']
    assert json.loads(job.import_sources) == source
    assert db.list_clients() == []
    attachments = db.list_attachments(job.id)
    assert len(attachments) == 1
    assert (db.config.ARCHIVE_ROOT/job.archive_path/'dokumenty'/attachments[0].stored_name).read_bytes() == columns_pdf.read_bytes()
    pane.load_job(job)
    assert pane.pod_requirements_edit.toPlainText() == values['pod_requirements']
    assert pane.settlement.fields['price'].text() == '650,00'
    assert pane.collect().local_only
    pane.close()


def test_incomplete_import_does_not_keep_today_and_remains_findable(app, local_db):
    from markus_dispatcher.ui.job_pane import JobPane
    pane = JobPane()
    pane.new_job(driver_id=local_db.id)
    pane._apply_document_import({'ref_numbers':'SYNTH-INCOMPLETE'}, {})
    job = pane.collect()
    assert job.load_datetime == job.unload_datetime == ''
    assert 'termin' in pane.import_hint.text()
    saved = db.save_job(job)
    pane.bind_driver(local_db.id)
    assert pane._job_id == saved.id
    assert pane.history_combo.findData(saved.id) >= 0
    assert pane.history_combo.itemText(pane.history_combo.findData(saved.id)).startswith('⚠')
    pane.close()


def test_local_only_is_sticky_and_remote_cannot_overwrite(local_db):
    job = db.save_job(JobData(driver_id=local_db.id, local_only=True, pod_requirements='synthetic POD'))
    job.local_only = False
    assert db.save_job(job).local_only
    remote = db.apply_remote_job({'ref_numbers':'remote'}, local_id=job.id)
    assert remote.local_only and remote.pod_requirements == 'synthetic POD'
    assert remote.ref_numbers != 'remote'


def test_database_migration_preserves_existing_jobs_and_adds_local_fields(local_db):
    original = db.save_job(JobData(driver_id=local_db.id, ref_numbers='SYNTH-OLD'))
    with db.get_conn() as connection:
        for column in ('local_only', 'pod_requirements', 'import_sources'):
            connection.execute(f'ALTER TABLE jobs DROP COLUMN {column}')
    db.init_db()
    restored = db.get_job(original.id)
    assert restored.ref_numbers == 'SYNTH-OLD'
    assert not restored.local_only and restored.pod_requirements == restored.import_sources == ''


def test_ui_blocks_external_actions_on_local_import(app, local_db, monkeypatch):
    from markus_dispatcher.ui import job_pane as module
    from markus_dispatcher.ui.job_pane import JobPane
    monkeypatch.setattr(module.QMessageBox, 'information', lambda *a, **k: None)
    monkeypatch.setattr(module, 'open_web_url', lambda *a, **k: pytest.fail('No external browser'))
    monkeypatch.setattr(module, 'start_geocode', lambda *a, **k: pytest.fail('No geocoding'))
    pane = JobPane()
    pane.new_job(driver_id=local_db.id)
    pane._apply_document_import({'client_name':'Synthetic client'}, {})
    pane._send_email()
    pane._send_whatsapp()
    pane._request_geocode('load')
    pane._open_maps()
    with pytest.raises(ValueError):
        pane.settlement._publish()
    assert db.list_clients() == []
    pane.close()


def test_firebase_and_google_refuse_imported_job_before_any_request(local_db, monkeypatch, tmp_path):
    from markus_dispatcher.integrations import firebase_sync
    from markus_dispatcher.core.google_sheets import Sheets
    job = db.save_job(JobData(driver_id=local_db.id, local_only=True, archive_path='synthetic-job'))
    class Collection:
        def document(self, *args):
            pytest.fail('Imported job must not access remote document')
    root = SimpleNamespace(collection=lambda *args: Collection())
    assert firebase_sync._push_jobs(root) == 0
    monkeypatch.setattr(Sheets, 'request', lambda *a, **k: pytest.fail('No Google request'))
    with pytest.raises(ValueError):
        Sheets().publish(job.id, 'file', 0, [])
    with pytest.raises(ValueError):
        firebase_sync._job_doc(job)
    bucket = SimpleNamespace(name='synthetic', blob=lambda *args: pytest.fail('No upload'))
    monkeypatch.setattr(firebase_sync, '_open_storage_bucket', lambda *args: bucket)
    monkeypatch.setattr(firebase_sync, '_remember_bucket', lambda *args: None)
    documents = db.config.ARCHIVE_ROOT/job.archive_path/'dokumenty'
    documents.mkdir(parents=True)
    (documents/'synthetic.pdf').write_bytes(b'fake')
    assert firebase_sync._sync_files(root, 'synthetic-user', None) == (0,0)


def test_clipboard_image_is_a_local_source_and_cancel_does_not_save(app, local_db):
    from PySide6.QtGui import QImage, QGuiApplication
    from markus_dispatcher.ui.document_import_dialog import DocumentImportDialog
    image = QImage(300, 400, QImage.Format.Format_RGB32)
    image.fill(0xffffff)
    previous = QGuiApplication.clipboard().image()
    try:
        QGuiApplication.clipboard().setImage(image)
        dialog = DocumentImportDialog()
        dialog.paste_image()
        assert dialog._has_document and dialog.source_path is None
        assert dialog.clipboard_image.size() == image.size()
        dialog.reject()
        assert db.list_jobs(include_done=True) == []
    finally:
        QGuiApplication.clipboard().setImage(previous)
