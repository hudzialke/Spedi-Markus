"""Large previews remain complete while each Qt table has a bounded size."""
from collections import Counter
import csv
from html.parser import HTMLParser
from io import StringIO

from markus_dispatcher.core import document_import as readers


class Cells(HTMLParser):
    def __init__(self, html):
        super().__init__()
        self.values = []
        self.spans = []
        self.current = None
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        if tag == 'td':
            self.current = []
            self.spans.append(dict(attrs))
        elif tag == 'br' and self.current is not None:
            self.current.append('\n')

    def handle_data(self, data):
        if self.current is not None:
            self.current.append(data)

    def handle_endtag(self, tag):
        if tag == 'td':
            self.values.append(''.join(self.current))
            self.current = None


def test_large_csv_preserves_every_value_across_pages(tmp_path):
    path = tmp_path / 'generated.csv'
    rows = [[f'Cell {r}/{c} <&>' for c in range(10)] for r in range(1000)]
    with path.open('w', encoding='utf-8', newline='') as output:
        csv.writer(output, delimiter=';').writerows(rows)
    sections = readers.read_spreadsheet(path)
    assert len(sections) == 10
    assert sections[0].title == 'CSV · wiersze 1–100 · kolumny 1–10'
    assert sections[-1].title == 'CSV · wiersze 901–1000 · kolumny 1–10'
    parsed = [Cells(section.html) for section in sections]
    assert all(len(page.values) <= readers.PREVIEW_PAGE_CELLS for page in parsed)
    assert Counter(value for page in parsed for value in page.values) == Counter(value for row in rows for value in row)


def test_csv_keeps_quoted_multiline_values(tmp_path):
    path = tmp_path / 'generated.csv'
    stream = StringIO(newline='')
    csv.writer(stream, delimiter=';').writerows([['LOAD', 'UNLOAD'], ['Street 1\nTown <A>', 'Other; road']])
    path.write_text(stream.getvalue(), encoding='utf-8', newline='')
    sections = readers.read_spreadsheet(path)
    assert len(sections) == 1 and sections[0].title == 'CSV'
    assert Cells(sections[0].html).values == ['LOAD', 'UNLOAD', 'Street 1\nTown <A>', 'Other; road']


def test_very_wide_ragged_csv_does_not_expand_to_a_dense_rectangle(tmp_path):
    path = tmp_path / 'generated.csv'
    rows = [[f'Wide {c}' for c in range(101)], ['Short'], *([[]] * 99)]
    with path.open('w', encoding='utf-8', newline='') as output:
        csv.writer(output, delimiter=';').writerows(rows)
    sections = readers.read_spreadsheet(path)
    assert len(sections) == 4
    assert sum(len(Cells(section.html).values) for section in sections) == 102
    assert 'kolumny 81–101' in sections[2].title


def test_xlsx_preserves_values_and_clips_merges_at_page_boundaries(tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = 'Synthetic'
    for row in range(1, 111):
        for col in range(1, 46):
            sheet.cell(row, col, f'R{row}C{col}')
    sheet.merge_cells('AL48:AR52')
    sheet['AL48'] = 'MERGED <VALUE>'
    expected = Counter(cell.value for row in sheet for cell in row if cell.value is not None)
    path = tmp_path / 'generated.xlsx'
    book.save(path)
    sections = readers.read_spreadsheet(path)
    assert len(sections) == 6
    parsed = [Cells(section.html) for section in sections]
    actual = Counter(value for page in parsed for value in page.values if value)
    expected['MERGED <VALUE>'] = 4
    assert actual == expected
    assert all(int(span.get('rowspan', 1)) <= 50 and int(span.get('colspan', 1)) <= 40
               for page in parsed for span in page.spans)
    assert sections[-1].title == 'Synthetic · wiersze 101–110 · kolumny 41–45'


def test_last_preview_page_can_assign_text_and_retains_source_range(tmp_path, monkeypatch, wait_for):
    monkeypatch.setenv('QT_QPA_PLATFORM', 'offscreen')
    monkeypatch.setenv('TEMP', str(tmp_path))
    monkeypatch.setenv('TMP', str(tmp_path))
    from PySide6.QtWidgets import QApplication
    from markus_dispatcher.ui.document_import_dialog import DocumentImportDialog
    app = QApplication.instance() or QApplication([])
    path = tmp_path / 'generated.csv'
    with path.open('w', encoding='utf-8', newline='') as output:
        csv.writer(output, delimiter=';').writerows([[f'R{r}C{c}' for c in range(10)] for r in range(301)])
    dialog = DocumentImportDialog()
    dialog.open_path(path)
    wait_for(lambda: dialog._reader is None)
    assert dialog.page_combo.count() == 4
    dialog.page_combo.setCurrentIndex(3)
    cursor = dialog.text_preview.document().find('R300C9')
    assert not cursor.isNull()
    dialog.text_preview.setTextCursor(cursor)
    dialog.assign_selection()
    assert dialog.values['ref_numbers'] == 'R300C9'
    assert dialog.sources['ref_numbers']['section'] == 'CSV · wiersze 301–301 · kolumny 1–10'
    dialog.page_combo.setCurrentIndex(0)
    assert dialog.selection.toPlainText() == ''
    dialog.reject()
    dialog.deleteLater()
    app.processEvents()
